import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from backend.app.config import settings
from backend.app.models.domain import Application, User, UserProfile, FinancialProfile, ExistingLoan, LoanRequest, MLPrediction, EligibilityResult, Consent, Bank, LoanProduct

class ExcelExporterService:
    @staticmethod
    def mask_mobile(mobile: str) -> str:
        if not mobile or len(mobile) < 7:
            return mobile or ""
        return mobile[:3] + "*****" + mobile[-2:]

    @staticmethod
    def mask_pan(pan: str) -> str:
        if not pan or len(pan) != 10:
            return pan or ""
        return pan[:3] + "*****" + pan[-1:]

    @classmethod
    def generate_applications_report(cls, db: Session, is_admin: bool = True) -> str:
        wb = Workbook()
        
        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        sub_font = Font(name="Calibri", size=10, italic=True, color="475569")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        
        # Remove default sheet
        wb.remove(wb.active)
        
        # 1. Applications Sheet
        ws_app = wb.create_sheet(title="Applications")
        ws_app.append(["LoanCompare AI - Applications Master Report"])
        ws_app.append(["Generated on: " + os.getenv("CURRENT_TIME", "2026-09-22") + " | Mode: Demo & Academic Audit"])
        ws_app.append([])
        headers_app = ["Application ID", "User Name", "Mobile (Masked)", "Email", "Status", "Risk Category", "ML Score (%)", "Created At"]
        ws_app.append(headers_app)
        
        apps = db.query(Application).order_by(Application.created_at.desc()).all()
        for app in apps:
            u = app.user
            ws_app.append([
                app.id,
                u.full_name if u else "N/A",
                cls.mask_mobile(u.mobile_number) if u else "N/A",
                u.email if u else "N/A",
                app.status,
                app.risk_category or "PENDING",
                f"{app.ml_score:.1f}" if app.ml_score is not None else "N/A",
                app.created_at.strftime("%Y-%m-%d %H:%M:%S")
            ])
            
        # 2. Applicants Sheet
        ws_user = wb.create_sheet(title="Applicants")
        ws_user.append(["User ID", "Full Name", "Mobile", "Email", "Age", "Employment Type", "Monthly Income (INR)", "Dependents"])
        users = db.query(User).all()
        for u in users:
            prof = u.profile
            ws_user.append([
                u.id,
                u.full_name,
                cls.mask_mobile(u.mobile_number),
                u.email,
                prof.age if prof else "N/A",
                prof.employment_type if prof else "N/A",
                prof.monthly_income if prof else 0.0,
                prof.dependents if prof else 0
            ])

        # 3. Financial Profile Sheet
        ws_fin = wb.create_sheet(title="Financial Profile")
        ws_fin.append(["Application ID", "Total Monthly Income", "Existing EMI Total", "Credit Card Outstanding", "Credit History", "Simulated Credit Score", "DTI Ratio (%)", "LTI Ratio"])
        fins = db.query(FinancialProfile).all()
        for f in fins:
            ws_fin.append([
                f.application_id,
                f.total_monthly_income,
                f.existing_emi_total,
                f.credit_card_outstanding,
                f.credit_history,
                f.credit_score_simulated or "N/A",
                f.dti_ratio,
                f.lti_ratio
            ])

        # 4. Existing Loans Sheet
        ws_loan = wb.create_sheet(title="Existing Loans")
        ws_loan.append(["Application ID", "Loan Type", "Original Amount", "Outstanding Amount", "Monthly EMI", "Remaining Tenure (Months)"])
        ex_loans = db.query(ExistingLoan).all()
        for el in ex_loans:
            ws_loan.append([
                el.application_id,
                el.loan_type,
                el.original_amount,
                el.outstanding_amount,
                el.monthly_emi,
                el.remaining_tenure_months
            ])

        # 5. ML Results Sheet
        ws_ml = wb.create_sheet(title="ML Results")
        ws_ml.append(["Application ID", "Risk Category", "Probability Score (%)", "Confidence (%)", "Model Version", "Prediction Time"])
        preds = db.query(MLPrediction).all()
        for p in preds:
            ws_ml.append([
                p.application_id,
                p.risk_category,
                p.probability_score,
                p.confidence,
                p.model_version,
                p.prediction_timestamp.strftime("%Y-%m-%d %H:%M:%S")
            ])

        # 6. Bank Matching Sheet
        ws_match = wb.create_sheet(title="Bank Matching")
        ws_match.append(["Application ID", "Bank & Product", "Match Status", "Matched Criteria", "Total Criteria", "Created At"])
        res_list = db.query(EligibilityResult).all()
        for r in res_list:
            prod = r.product
            bank_p_name = f"{prod.bank.name} - {prod.product_name}" if prod and prod.bank else "Product ID " + str(r.product_id)
            ws_match.append([
                r.application_id,
                bank_p_name,
                r.match_status,
                r.matched_criteria_count,
                r.total_criteria_count,
                r.created_at.strftime("%Y-%m-%d %H:%M:%S")
            ])

        # 7. Consent Logs Sheet
        ws_consent = wb.create_sheet(title="Consent Logs")
        ws_consent.append(["Consent ID", "User ID", "Consent Given", "Purpose", "Consent Version", "IP Address", "Timestamp"])
        consents = db.query(Consent).all()
        for c in consents:
            ws_consent.append([
                c.id,
                c.user_id,
                "YES" if c.consent_given else "NO",
                c.purpose,
                c.consent_version,
                c.ip_address or "127.0.0.1",
                c.created_at.strftime("%Y-%m-%d %H:%M:%S")
            ])

        # Styling formatting pass across all sheets
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    cell.border = thin_border
            # Format header row
            if sheet.max_row >= 1:
                header_row_idx = 4 if sheet.title == "Applications" else 1
                for cell in sheet[header_row_idx]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Auto-adjust column width
            for col in sheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

        file_path = os.path.join(settings.EXPORTS_DIR, "loan_applications.xlsx")
        wb.save(file_path)
        return file_path
